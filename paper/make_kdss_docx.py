# -*- coding: utf-8 -*-
"""Pour paper/draft_kdss.html into the KDSS 2026 fall conference template.

Writes paper/draft_kdss.docx. Each element of the manuscript is cloned from the template's own sample of that element
(title, author, abstract, section heading, body, figure, caption, table, references), so fonts and the font sizes the
template prescribes (title 14, abstract heading 10, abstract and body 9, section headings 11, captions 9) are kept.
Margins, line spacing, paragraph spacing and figure sizes are tightened (LAYOUT below) to fit the 2-page limit for
oral/poster papers; the references may run past page 2. draft_kdss.html is draft.html shortened for this venue.

Usage: python make_kdss_docx.py   (needs python-docx; imports the HTML parser from make_docx.py)
"""
import copy
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.text.paragraph import Paragraph

import make_docx

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "한국데이터사이언스_2026추계학술대회_논문양식[논문].docx"
SRC, OUT = HERE / "draft_kdss.html", HERE / "draft_kdss.docx"

TITLE_KO = "결정 모델은 VLM 로봇 제어의 빠른 계층이 될 수 있는가: RoboTwin 2.0 실패 분석과 피지컬 AI에서의 한계"
TITLE_EN = ("Can Decision Models Serve as the Fast Layer of VLM-Based Robot Control? "
            "Failure Analysis on RoboTwin 2.0 and Limitations in Physical AI")
AUTHORS = [("임준현", "*"), (", 성낙준", "**")]
AFFILS = [("*", "한신대학교 컴퓨터공학부"), ("**", "한신대학교 AI·SW학")]
EMAILS = "tomtom35177@hs.ac.kr, njsung@hs.ac.kr"

TW = 56.7                        # twips per mm
# page: the template has 30/25/25/25 mm margins; header and footer text are the template's own
MARGIN = dict(top=17, bottom=15, left=15, right=15, header=8, footer=6)  # mm
COL_GAP_MM = 5
TEXT_W_MM = 210 - MARGIN["left"] - MARGIN["right"]
COL_W_MM = (TEXT_W_MM - COL_GAP_MM) / 2
FIG1_W_MM = 158
TABLE_WIDTHS = [13, 23, 13, 33, 18]  # percent of the column, table 1
# exact line heights in twips (template: 160% of the font size); sizes are the template's
LINE = dict(title=330, author=230, abs_head=270, abstract=240, head=270, body=236, caption=222, cell=186,
            tnote=176, ref=228)
TABLE_PT, TNOTE_PT = 8, 7.5      # not prescribed by the template (its sample table is 9 pt)
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
BODY_INDENT = 180                # twips, first-line indent of a body paragraph that follows another one
HEAD_BEFORE, HEAD_AFTER = 70, 20  # twips, around section headings II.. (the template uses blank lines)


def text_of(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def strip_ids(el):
    for e in el.iter():
        for a in list(e.attrib):
            if a.endswith("}paraId") or a.endswith("}textId"):
                del e.attrib[a]
    return el


def proto_rpr(p):
    r = p.find(qn("w:r"))
    rpr = r.find(qn("w:rPr")) if r is not None else None
    return copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")


def make_run(rpr, text, bold=False, italic=False, sup=False, size=None):
    r = OxmlElement("w:r")
    rp = copy.deepcopy(rpr)
    for tag in ("w:vertAlign",):
        for e in rp.findall(qn(tag)):
            rp.remove(e)
    if bold:
        rp.get_or_add_b()
        if rp.find(qn("w:bCs")) is None:
            rp.find(qn("w:b")).addnext(OxmlElement("w:bCs"))
    if italic:
        rp.get_or_add_i()
    if size is not None:
        for tag in ("w:sz", "w:szCs"):
            for e in rp.findall(qn(tag)):
                rp.remove(e)
        rp.get_or_add_sz().set(qn("w:val"), str(int(size * 2)))
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rp.find(qn("w:sz")).addnext(szcs)
    if sup:
        rp.get_or_add_vertAlign().set(qn("w:val"), "superscript")
    r.append(rp)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def fill(p, runs, rpr=None, **kw):
    """Replace a paragraph's content with runs [(text, bold, italic)], keeping its pPr and first run format."""
    rpr = proto_rpr(p) if rpr is None else rpr
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    for text, b, i in runs:
        parts = text.split("\n")
        for k, part in enumerate(parts):
            if k:
                br = OxmlElement("w:r")
                br.append(copy.deepcopy(rpr))
                br.append(OxmlElement("w:br"))
                p.append(br)
            if part:
                p.append(make_run(rpr, part, bold=b, italic=i, **kw))
    return p


def clone(proto):
    p = strip_ids(copy.deepcopy(proto))
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    return p


def set_spacing(p, **attrs):
    ppr = p.get_or_add_pPr()
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.insert_element_before(sp, "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
                                  "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    for k, v in attrs.items():
        sp.set(qn(f"w:{k}"), str(v))


def exact(p, line, before=0, after=0):
    set_spacing(p, before=before, after=after, line=line, lineRule="exact")


def set_page(sect):
    mar = sect.find(qn("w:pgMar"))
    for k, mm in MARGIN.items():
        mar.set(qn(f"w:{k}"), str(round(mm * TW)))
    cols = sect.find(qn("w:cols"))
    if cols is not None:
        cols.set(qn("w:space"), str(round(COL_GAP_MM * TW)))


def set_ind(p, **attrs):
    ppr = p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.insert_element_before(ind, "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                                  "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                                  "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    for k, v in attrs.items():
        ind.set(qn(f"w:{k}"), str(v))


def set_jc(p, val):
    ppr = p.get_or_add_pPr()
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        ppr.insert_element_before(jc, "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                                  "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    jc.set(qn("w:val"), val)


def keep_next(p):
    ppr = p.get_or_add_pPr()
    if ppr.find(qn("w:keepNext")) is None:
        ppr.insert_element_before(OxmlElement("w:keepNext"), *[t for t in (
            "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
            "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
            "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
            "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
            "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
            "w:rPr", "w:sectPr", "w:pPrChange")])


def keep_lines(p):
    ppr = p.get_or_add_pPr()
    if ppr.find(qn("w:keepLines")) is None:
        kl = OxmlElement("w:keepLines")
        kn = ppr.find(qn("w:keepNext"))
        if kn is not None:
            kn.addnext(kl)
        else:
            ppr.insert(1 if ppr.find(qn("w:pStyle")) is not None else 0, kl)


def wrap_words(p):
    ppr = p.get_or_add_pPr()
    for e in ppr.findall(qn("w:wordWrap")):
        ppr.remove(e)
    ppr.insert_element_before(OxmlElement("w:wordWrap"), "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                              "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
                              "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                              "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                              "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def roman_heading(runs):
    text = "".join(t for t, _, _ in runs)
    m = re.match(r"(\d+)\.\s*(.*)", text)
    return f"{ROMAN[int(m.group(1)) - 1]}. {m.group(2)}" if m else text


def build():
    html = SRC.read_text(encoding="utf-8")
    parser = make_docx.Parser()
    parser.feed(html)
    blocks = parser.blocks

    doc = Document(TEMPLATE)
    body = doc.element.body
    final_sect = body.find(qn("w:sectPr"))

    # ---- template samples (taken before the sample content is removed) ----
    paras = [c for c in body if c.tag == qn("w:p")]
    brk = next(p for p in paras if p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
    after = list(body)[list(body).index(brk) + 1:]
    after = [c for c in after if c is not final_sect]

    def sample(pred):
        return next(c for c in after if c.tag == qn("w:p") and pred(c))

    P_HEAD = sample(lambda p: text_of(p).startswith("I. 서론"))
    P_BODY = sample(lambda p: text_of(p).startswith("데이터사이언스는"))
    P_FIG = sample(lambda p: p.find(".//" + qn("w:drawing")) is not None)
    P_CAP = sample(lambda p: text_of(p).startswith("그림 1."))
    P_TCAP = sample(lambda p: text_of(p).startswith("표1."))
    P_REFHEAD = sample(lambda p: text_of(p).startswith("참고문헌"))
    P_REF = sample(lambda p: text_of(p).startswith("[2]"))
    T_SAMPLE = next(c for c in after if c.tag == qn("w:tbl"))
    P_TCELL = T_SAMPLE.find(".//" + qn("w:p"))
    rpr_body = proto_rpr(P_BODY)
    rpr_cap = proto_rpr(P_CAP)
    rpr_head = proto_rpr(P_HEAD)
    rpr_ref = proto_rpr(P_REF)
    rpr_tcap = proto_rpr(P_TCAP)
    rpr_tcell = proto_rpr(P_TCELL)
    P_HEAD, P_BODY, P_FIG, P_CAP, P_TCAP, P_REFHEAD, P_REF, P_TCELL = (
        copy.deepcopy(x) for x in (P_HEAD, P_BODY, P_FIG, P_CAP, P_TCAP, P_REFHEAD, P_REF, P_TCELL))
    for c in after:
        body.remove(c)

    # ---- page setup for both sections; header and footer paragraphs lose the default 8 pt space after ----
    set_page(final_sect)
    set_page(brk.find(qn("w:pPr")).find(qn("w:sectPr")))
    for part in (doc.sections[0].header, doc.sections[0].footer):
        for hp in part.paragraphs:
            set_spacing(hp._p, before=0, after=0, line=240, lineRule="auto")

    # ---- the instruction box "구두/포스터 발표 논문: 2페이지 이내" sits in the section-break paragraph ----
    for child in list(brk):
        if child.tag != qn("w:pPr"):
            brk.remove(child)
    set_spacing(brk, before=0, after=0, line=240, lineRule="auto")
    brk_rpr = brk.find(qn("w:pPr")).find(qn("w:rPr"))
    if brk_rpr is None:
        brk_rpr = OxmlElement("w:rPr")
        brk.find(qn("w:pPr")).insert_element_before(brk_rpr, "w:sectPr", "w:pPrChange")
    for e in brk_rpr.findall(qn("w:sz")) + brk_rpr.findall(qn("w:szCs")):
        brk_rpr.remove(e)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    brk_rpr.append(sz)

    # ---- title block (the template's table) ----
    tb = doc.tables[0]
    side = 60
    widths = [side, round(TEXT_W_MM * TW) - 2 * side, side]
    for gc, w in zip(tb._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for row in tb.rows:
        trpr = row._tr.trPr
        if trpr is not None:
            for h in trpr.findall(qn("w:trHeight")):
                trpr.remove(h)
        for c, cell in enumerate(row.cells):
            cell.width = Mm(widths[c] / TW)
            for cp_ in cell.paragraphs:
                exact(cp_._p, 120)
    tb._tbl.remove(tb.rows[1]._tr)  # empty spacer row between title and authors
    ko, en = tb.cell(0, 1).paragraphs[:2]
    fill(ko._p, [(TITLE_KO, False, False)])
    fill(en._p, [(TITLE_EN, False, False)])
    exact(ko._p, LINE["title"])
    exact(en._p, LINE["title"], before=40)
    au, af, em = tb.cell(1, 1).paragraphs[:3]
    exact(au._p, LINE["author"], before=140)
    exact(af._p, LINE["author"])
    exact(em._p, LINE["author"])
    rpr_au = proto_rpr(au._p)
    for child in list(au._p):
        if child.tag != qn("w:pPr"):
            au._p.remove(child)
    for name, mark in AUTHORS:
        au._p.append(make_run(rpr_au, name))
        au._p.append(make_run(rpr_au, mark, sup=True))
    rpr_af = proto_rpr(af._p)
    for child in list(af._p):
        if child.tag != qn("w:pPr"):
            af._p.remove(child)
    for k, (mark, unit) in enumerate(AFFILS):
        af._p.append(make_run(rpr_af, mark, sup=True))
        af._p.append(make_run(rpr_af, unit + (", " if k < len(AFFILS) - 1 else "")))
    fill(em._p, [(EMAILS, False, False)])
    ab_cell = tb.cell(2, 1)
    ab_head, ab_text = ab_cell.paragraphs[:2]
    fill(ab_head._p, [("요 약", False, False)])
    exact(ab_head._p, LINE["abs_head"], before=120, after=20)
    exact(ab_text._p, LINE["abstract"], after=60)
    abstract = next(b for b in blocks if b["kind"] == "abstract")
    runs = make_docx.clean_runs(abstract["runs"])
    if runs and runs[0][0].strip() == "요약":  # the HTML starts the abstract with a bold "요약" label
        runs = runs[1:]
    runs = make_docx.clean_runs(runs)
    fill(ab_text._p, runs)
    for extra in ab_cell.paragraphs[2:]:
        if not extra.text.strip():
            extra._p.getparent().remove(extra._p)

    # ---- body ----
    def add_before_final(el):
        final_sect.addprevious(el)
        return el

    def add_picture(container_add, path, width_mm):
        p = clone(P_FIG)
        set_spacing(p, before=40, after=0, line=240, lineRule="auto")
        set_ind(p, firstLine=0)
        keep_next(p)
        container_add(p)
        Paragraph(p, doc._body).add_run().add_picture(str(path), width=Mm(width_mm))
        return p

    def add_caption(container_add, runs):
        runs = [(t, False, i) for t, b, i in runs]  # template captions are not bold
        p = fill(clone(P_CAP), runs, rpr=rpr_cap)
        exact(p, LINE["caption"], before=30, after=80)
        set_jc(p, "both")
        keep_lines(p)
        return container_add(p)

    pending_wide_caption = False
    first_head = True
    prev = None
    for blk in blocks:
        kind = blk["kind"]
        if kind in ("h1", "authors", "affil", "abstract", "cols_start"):
            continue
        if kind == "img" and blk["wide"]:
            add_picture(brk.addprevious, HERE / blk["src"], FIG1_W_MM)
            pending_wide_caption = True
        elif kind == "figcaption" and pending_wide_caption:
            add_caption(brk.addprevious, make_docx.clean_runs(blk["runs"]))
            pending_wide_caption = False
        elif kind == "img":
            add_picture(add_before_final, HERE / blk["src"], COL_W_MM)
        elif kind == "figcaption":
            add_caption(add_before_final, make_docx.clean_runs(blk["runs"]))
        elif kind == "h2":
            p = fill(clone(P_HEAD), [(roman_heading(blk["runs"]), False, False)], rpr=rpr_head)
            exact(p, LINE["head"], before=0 if first_head else HEAD_BEFORE, after=HEAD_AFTER)
            keep_next(p)
            first_head = False
            add_before_final(p)
        elif kind == "p":
            p = fill(clone(P_BODY), make_docx.clean_runs(blk["runs"]), rpr=rpr_body)
            exact(p, LINE["body"])
            if prev == "p":
                set_ind(p, firstLine=BODY_INDENT)
            add_before_final(p)
        elif kind == "tcap":
            p = fill(clone(P_TCAP), [(t, False, i) for t, b, i in make_docx.clean_runs(blk["runs"])], rpr=rpr_tcap)
            exact(p, LINE["body"], before=60, after=20)
            keep_next(p)
            add_before_final(p)
        elif kind == "table":
            add_table(doc, blk["rows"], P_TCELL, rpr_tcell, add_before_final)
        elif kind == "tnote":
            p = fill(clone(P_CAP), make_docx.clean_runs(blk["runs"]), rpr=rpr_cap, size=TNOTE_PT)
            set_jc(p, "left")  # justified 7.5 pt notes stretch badly with whole-word wrapping
            exact(p, LINE["tnote"], before=30, after=80)
            add_before_final(p)
        elif kind == "refs_h2":
            p = fill(clone(P_REFHEAD), [("참고문헌", False, False)], rpr=proto_rpr(P_REFHEAD))
            exact(p, LINE["head"], before=HEAD_BEFORE, after=HEAD_AFTER)
            keep_next(p)
            add_before_final(p)
        elif kind == "ref":
            p = fill(clone(P_REF), make_docx.clean_runs(blk["runs"]), rpr=rpr_ref)
            exact(p, LINE["ref"])
            add_before_final(p)
        else:
            raise ValueError(f"unhandled block {kind}")
        prev = kind

    cp = doc.core_properties
    cp.title = TITLE_KO
    cp.author = "임준현"
    cp.last_modified_by = "임준현"
    doc.save(OUT)
    print(OUT)


def add_table(doc, rows, p_cell, rpr_cell, add):
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = doc.styles["Table Grid"]
    tbl = t._tbl
    tbl.getparent().remove(tbl)
    add(tbl)
    col_tw = [int(COL_W_MM * 56.7 * w / 100) for w in TABLE_WIDTHS]
    tblpr = tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    tblw.set(qn("w:w"), str(sum(col_tw)))
    tblw.set(qn("w:type"), "dxa")
    for gc, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), col_tw):
        gc.set(qn("w:w"), str(w))
    mar = OxmlElement("w:tblCellMar")
    for edge, tw in (("top", 20), ("left", 57), ("bottom", 20), ("right", 57)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(tw))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.find(qn("w:tblLook")).addprevious(mar)
    occupied = set()
    for r, hrow in enumerate(rows):
        c = 0
        for hcell in hrow:
            while (r, c) in occupied:
                c += 1
            cell = t.cell(r, c)
            span = hcell["rowspan"]
            if span > 1:
                cell = cell.merge(t.cell(r + span - 1, c))
                occupied.update((r + k, c) for k in range(1, span))
            p = cell.paragraphs[0]._p
            new = fill(clone(p_cell), make_docx.clean_runs(hcell["runs"]), rpr=rpr_cell, size=TABLE_PT)
            if hcell["header"]:
                for rr in new.findall(qn("w:r")):
                    rp = rr.find(qn("w:rPr"))
                    rp.get_or_add_b()
            exact(new, LINE["cell"])
            wrap_words(new)  # "물체 모니터" breaks at the space, not inside a word
            set_jc(new, "left")
            keep_next(new)
            p.addprevious(new)
            p.getparent().remove(p)
            c += 1
    for r, row in enumerate(t.rows):
        for c, cell in enumerate(row.cells):
            cell.width = Mm(COL_W_MM * TABLE_WIDTHS[c] / 100)


if __name__ == "__main__":
    build()
